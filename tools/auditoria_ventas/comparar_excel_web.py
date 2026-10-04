"""
comparar_excel_web.py
---------------------
Compara, celda por celda, el Excel del Informe Diario GRIDO contra la pantalla
web (usp_InformeDiarioGrido) para una o varias jornadas. Las dos salidas se
calculan por caminos distintos (el Excel lee SmartFran directo, la web lee la
huella), asi que cualquier diferencia es un error en uno de los dos.

Uso (con el Python del entorno del Excel):
    python comparar_excel_web.py 2026-09-26 [2026-09-27 ...]
    python comparar_excel_web.py 2026-09-26 --carpeta <app del Excel>

Por defecto genera el Excel con la copia de desarrollo de este repositorio
(reporting/informe_grido) en modo --sin-mail --sin-drive, y lo compara.
Codigo de salida: 0 si no hay diferencias, 1 si las hay.

Creado: 2026-10-02
"""
import argparse
import os
import subprocess
import sys

import openpyxl
import pyodbc

REPO = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
CARPETA_DEV = os.path.join(REPO, "reporting", "informe_grido")
CONN = ("DRIVER={ODBC Driver 17 for SQL Server};SERVER=WIN-6ARG3SUELOE\\SQLEXPRESS;"
        "DATABASE=DF_DTW;Trusted_Connection=yes;TrustServerCertificate=yes")


def leer_excel(ruta):
    ws = openpyxl.load_workbook(ruta)["Informe Diario"]
    # Dos formatos de hoja:
    #  - octubre 2026 (espec de Damian): P Kilos Club Grido, T/U/V clima de zona
    #    con condicion, y clima tambien en subtotales y TOTAL.
    #  - anterior: P Ventas Club, Q Kilos Club, U/V clima por sucursal.
    nuevo = "Kilos" in str(ws.cell(9, 16).value) and "Grido" in str(ws.cell(9, 16).value)
    c = ({"kc": 16, "vc": None, "st": 20, "mm": 21, "cond": 22} if nuevo
         else {"kc": 17, "vc": 16, "st": 21, "mm": 22, "cond": None})
    filas, subt = {}, {}
    for r in range(10, ws.max_row + 1):
        turno = ws.cell(r, 2).value
        clima = (ws.cell(r, c["st"]).value, ws.cell(r, c["mm"]).value,
                 ws.cell(r, c["cond"]).value if c["cond"] else None)
        if turno == "Todos" or ws.cell(r, 1).value == "TOTAL":
            subt[str(ws.cell(r, 1).value).strip()] = clima
            continue
        if not isinstance(turno, int):
            continue
        # La caja de delivery dice "DELI" en lugar del numero: la fila se
        # identifica por turno y cajero (el turno ya es de una sola caja).
        caja_txt = str(ws.cell(r, 3).value).strip()
        cajero = str(ws.cell(r, 1).value).strip().lower()
        filas[(turno, cajero)] = {
            "caja": None if caja_txt == "DELI" else int(caja_txt),
            "ventas": ws.cell(r, 7).value, "kilos": ws.cell(r, 6).value,
            "tickets": ws.cell(r, 9).value, "promos": ws.cell(r, 13).value or 0,
            "ventas_club": ws.cell(r, c["vc"]).value if c["vc"] else None,
            "kilos_club": ws.cell(r, c["kc"]).value or 0,
            "delivery": caja_txt == "DELI",
            "sens": clima[0], "mm": clima[1], "cond": clima[2],
        }
    return filas, subt, nuevo


def comparar(fecha, carpeta):
    py = sys.executable
    r = subprocess.run([py, "informe_cli.py", "--fecha", fecha, "--sin-mail", "--sin-drive"],
                       cwd=carpeta, capture_output=True, text=True)
    if r.returncode != 0:
        print(f"{fecha}: no se pudo generar el Excel (codigo {r.returncode})")
        return 1
    excel, subt, nuevo = leer_excel(os.path.join(carpeta, "salidas", f"{fecha}_InformeDiarioGRIDO.xlsx"))

    cur = pyodbc.connect(CONN).cursor()
    cur.execute("EXEC dbo.usp_InformeDiarioGrido @FechaDesde=?, @FechaHasta=?", fecha, fecha)
    cols = [c[0] for c in cur.description]
    web = [dict(zip(cols, x)) for x in cur.fetchall()]

    difs = []
    vistos = set()
    vistos_sub = set()
    for w in web:
        clave = (w["Turno"], w["Cajero"].strip().lower())
        vistos.add(clave)
        e = excel.get(clave)
        if e is None:
            difs.append(f"turno {w['Turno']} caja {w['Caja']} {w['Cajero']}: esta en la web y no en el Excel")
            continue
        pares = {
            "ventas": (w["Ventas"], e["ventas"]), "kilos": (w["Kilos"], e["kilos"]),
            "tickets": (w["Tickets"], e["tickets"]), "promos": (w["Promos"], e["promos"]),
            "kilos_club": (w.get("KilosClub"), e["kilos_club"]),
            "delivery": (bool(w.get("EsCajaDelivery")), e["delivery"]),
            # Si no es delivery, el numero de caja tiene que coincidir.
            "caja": (w["Caja"], w["Caja"] if e["delivery"] else e["caja"]),
        }
        if nuevo:
            pares.update({"sens": (w.get("ClimaZonaSens"), e["sens"]), "mm": (w.get("ClimaZonaLluvia"), e["mm"]),
                          "condicion": (w.get("ClimaZonaCondicion"), e["cond"])})
        else:
            pares.update({"ventas_club": (w["VentasClub"], e["ventas_club"]),
                          "sens": (w.get("SensacionTermica"), e["sens"]), "mm": (w.get("LluviaMm"), e["mm"])})
        # Clima del subtotal de la sucursal y del TOTAL (solo formato nuevo).
        if nuevo:
            for nombre, clave, cols3 in ((w["SucursalRotulo"], w["SucursalRotulo"], ("SucClimaZonaSens", "SucClimaZonaLluvia", "SucClimaZonaCondicion")),
                                         ("TOTAL", "TOTAL", ("DiaClimaZonaSens", "DiaClimaZonaLluvia", "DiaClimaZonaCondicion"))):
                if (fecha, clave) in vistos_sub:
                    continue
                vistos_sub.add((fecha, clave))
                xs = subt.get(clave, (None, None, None))
                for k, col, b in zip(("sens", "mm", "condicion"), cols3, xs):
                    a = w.get(col)
                    igual = (a == b) if (a is None or b is None or isinstance(a, str)) else abs(float(a) - float(b)) < 0.051
                    if not igual:
                        difs.append(f"{nombre} (subtotal) {k}: web={a} excel={b}")
        for k, (a, b) in pares.items():
            if a is None or b is None or isinstance(a, (bool, str)):
                igual = (a == b) or (a is None and b is None)
            else:
                igual = abs(float(a) - float(b)) < 0.051
            if not igual:
                difs.append(f"turno {w['Turno']} caja {w['Caja']} {w['Cajero']} {k}: web={a} excel={b}")
    for clave in set(excel) - vistos:
        difs.append(f"turno {clave[0]} {clave[1]}: esta en el Excel y no en la web")

    print(f"{fecha}: {len(excel)} filas, {len(difs)} diferencia(s)")
    for d in difs:
        print("   ", d)
    return 1 if difs else 0


def main():
    p = argparse.ArgumentParser()
    p.add_argument("fechas", nargs="+")
    p.add_argument("--carpeta", default=CARPETA_DEV)
    a = p.parse_args()
    rc = 0
    for f in a.fechas:
        rc |= comparar(f, a.carpeta)
    return rc


if __name__ == "__main__":
    sys.exit(main())
