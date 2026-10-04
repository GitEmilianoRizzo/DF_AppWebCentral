"""
Hojas extra del Informe Diario GRIDO
====================================
Productos, Promociones, Sobreventas, Tickets y Efemerides: las 5 hojas que se
agregan despues de "Informe Diario". Especificacion completa de Damian en
ESPEC_Hojas_Damian_2026-10-01.md (secciones 3 a 6); su ejemplo terminado del
27/09/2026 esta en ESPEC_ejemplo_2026-09-27.xlsx.

Las consultas van contra el warehouse (DF_DTW.dbo.TRX_HUELLA_VENTA), no
contra SmartFran: la tarea de las 13:00 corre despues de la carga de la
huella de las 12:30, asi que el dato ya esta.

Uso:
    datos = leer_datos(conn, fecha)            # con la conexion abierta
    agregar_hojas(wb, datos, fecha)            # despues de construir_excel()

Si algo falla al leer o armar estas hojas, el informe sale igual con la hoja
"Informe Diario": son complementarias y no pueden frenar el envio.

Creado: 2026-10-04
"""
import csv
import os
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta

import pandas as pd
from openpyxl.chart import BarChart, DoughnutChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.formatting.rule import CellIsRule, DataBarRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_DWH = "DF_DTW"
BASE_ORIGEN = "SRV_GRIDO_ZSUR"
DIR_CALENDARIO = os.path.join(os.path.dirname(os.path.abspath(__file__)), "calendario")

# Orden y colores de sucursal: los mismos del Informe Diario.
SUCURSALES = ["Fiorito", "Escalada", "Lanus Oeste"]
COLOR_SUC = {
    "Fiorito": ("FF4A7C59", "FFEBF5EE"),
    "Escalada": ("FF2E4053", "FFEAF0FB"),
    "Lanus Oeste": ("FF784212", "FFFEF5E7"),
}
PALETA = ["FF784212", "FFC0392B", "FF1F618D", "FF4A7C59", "FFD68910",
          "FF7D3C98", "FF5D6D7E", "FFA04000", "FF117A65", "FF2E4053"]
BANNER, SUBT, SUBT_TXT, TOTAL = "FF1E2D3D", "FF2C3E50", "FFD5DBE1", "FF1A252F"
SV_COLOR, PROMO_COLOR = "FF6B4226", "FF1F618D"
UMBRAL_SV = 0.15
DIAS = ["LUNES", "MARTES", "MIERCOLES", "JUEVES", "VIERNES", "SABADO", "DOMINGO"]
DIAS_ES = ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes", "Sábado", "Domingo"]

F_CANT, F_KG, F_PESOS, F_PCT = "#,##0", "#,##0.0", "$#,##0", "0.0%"


# ============================================================================
# Lectura
# ============================================================================
def _filtro():
    sucs = ",".join(f"'{s}'" for s in SUCURSALES)
    return (f"FROM {BASE_DWH}.dbo.TRX_HUELLA_VENTA "
            f"WHERE BASE_ORIGEN = '{BASE_ORIGEN}' AND FECHA_OPERATIVA = ? "
            f"AND ES_ANULADA = 0 AND SUCURSAL_DESCRIP IN ({sucs})")


def leer_datos(conn, fecha):
    """Las consultas Q1 a Q6 de la espec (seccion 4). `fecha` es la jornada."""
    f = _filtro()
    q = {
        # Q1. Todo lo vendido, sin las cabeceras de promo/sobreventa (PROMO=2).
        "productos": f"""
            SELECT RTRIM(SUCURSAL_DESCRIP) AS Suc,
                   ISNULL(NULLIF(RTRIM(ART_GRUPO_DESCRIP), ''), '(sin rubro)') AS Rubro,
                   RTRIM(ART_DESCRIP) AS Producto,
                   SUM(CANTIDAD) AS Cantidad, SUM(KILOS) AS Kilos, SUM(IMPORTE) AS Venta
            {f} AND ISNULL(PROMO, 0) <> 2
            GROUP BY SUCURSAL_DESCRIP, ART_GRUPO_DESCRIP, ART_DESCRIP""",
        # Q2. Productos dentro de cada promocion.
        "promo_det": f"""
            SELECT RTRIM(SUCURSAL_DESCRIP) AS Suc, RTRIM(PROMOCION_DESCRIP) AS Promocion,
                   RTRIM(ART_DESCRIP) AS Producto,
                   SUM(CANTIDAD) AS Cantidad, SUM(KILOS) AS Kilos, SUM(IMPORTE) AS Venta,
                   SUM(DESCUENTOS) AS Descuento
            {f} AND PROMOCION_TIPO = 'PROMOCION' AND PROMO = 1
            GROUP BY SUCURSAL_DESCRIP, PROMOCION_DESCRIP, ART_DESCRIP""",
        # Q3. Usos: una cabecera (PROMO=2) por cada vez que se aplico.
        "promo_usos": f"""
            SELECT RTRIM(SUCURSAL_DESCRIP) AS Suc, RTRIM(PROMOCION_DESCRIP) AS Promocion,
                   SUM(CASE WHEN PROMO = 2 THEN 1 ELSE 0 END) AS Usos
            {f} AND PROMOCION_TIPO = 'PROMOCION'
            GROUP BY SUCURSAL_DESCRIP, PROMOCION_DESCRIP""",
        # Q4. Productos aceptados como sobreventa.
        "sv_det": f"""
            SELECT RTRIM(SUCURSAL_DESCRIP) AS Suc, RTRIM(PROMOCION_DESCRIP) AS Sobreventa,
                   RTRIM(ART_DESCRIP) AS Producto,
                   SUM(CANTIDAD) AS Cantidad, SUM(KILOS) AS Kilos, SUM(IMPORTE) AS Venta,
                   SUM(DESCUENTOS) AS Descuento
            {f} AND PROMOCION_TIPO = 'SOBREVENTA' AND PROMO = 1
            GROUP BY SUCURSAL_DESCRIP, PROMOCION_DESCRIP, ART_DESCRIP""",
        # Q6. Un renglon por ticket.
        "tickets": f"""
            WITH l AS (
              SELECT RTRIM(SUCURSAL_DESCRIP) suc, RTRIM(USULOGIN) u, TICKET_KEY tk, FECHA_HORA fh, HORA,
                     SOBREVENTA sv, PROMOCION_TIPO pt, PROMO, ART_GRUPO_DESCRIP grp, ART_DESCRIP art,
                     CANTIDAD, IMPORTE, VTAIMPORTE, ES_PLATAFORMA_DELIVERY plat, PLATAFORMA
              {f}),
            t AS (SELECT suc, u, tk, MIN(HORA) hora, MIN(fh) fh, MAX(sv) sv, MAX(CAST(plat AS int)) plat,
                         MAX(ISNULL(PLATAFORMA, '')) plataforma, MAX(VTAIMPORTE) importe
                  FROM l GROUP BY suc, u, tk),
            mainp AS (SELECT tk, grp, ROW_NUMBER() OVER (PARTITION BY tk ORDER BY IMPORTE DESC, CANTIDAD DESC) rn
                      FROM l WHERE ISNULL(pt, '') <> 'SOBREVENTA' AND ISNULL(PROMO, 0) <> 2),
            prod AS (SELECT tk, STRING_AGG(CAST(CONCAT(CAST(CAST(CANTIDAD AS decimal(9,0)) AS varchar(10)), ' ',
                                LTRIM(RTRIM(art))) AS nvarchar(max)), ' + ') WITHIN GROUP (ORDER BY IMPORTE DESC) productos
                     FROM l WHERE ISNULL(PROMO, 0) <> 2 AND ISNULL(pt, '') <> 'SOBREVENTA' GROUP BY tk),
            psv AS (SELECT tk, STRING_AGG(CAST(CONCAT(CAST(CAST(CANTIDAD AS decimal(9,0)) AS varchar(10)), ' ',
                               LTRIM(RTRIM(art))) AS nvarchar(max)), ' + ') WITHIN GROUP (ORDER BY IMPORTE DESC) productos_sv
                    FROM l WHERE pt = 'SOBREVENTA' AND PROMO = 1 GROUP BY tk)
            SELECT t.suc AS Sucursal, t.u AS Cajero, t.hora AS Hora, t.fh AS FechaHora, t.tk AS Ticket,
                   CASE WHEN t.plat = 0 THEN 'Salón' WHEN t.plataforma = 'PosLocal' THEN 'Delivery propio'
                        WHEN t.plataforma <> '' THEN t.plataforma ELSE 'Delivery' END AS Canal,
                   t.importe AS Importe, ISNULL(LTRIM(RTRIM(m.grp)), '(sin dato)') AS RubroPrincipal,
                   p.productos AS ProductoLlevado,
                   CASE WHEN t.sv = 'ACEPTADA' THEN 'Aceptada' WHEN t.sv = 'RECHAZADA' THEN 'Rechazada'
                        ELSE 'No ofrecida' END AS Sobreventa,
                   s.productos_sv AS ProductoLlevadoEnSV
            FROM t LEFT JOIN mainp m ON m.tk = t.tk AND m.rn = 1
                   LEFT JOIN prod p ON p.tk = t.tk
                   LEFT JOIN psv s ON s.tk = t.tk""",
    }
    return {k: pd.read_sql(sql, conn, params=[fecha]) for k, sql in q.items()}


def _orden_suc(df, col="Suc"):
    return df.assign(_o=df[col].map({s: i for i, s in enumerate(SUCURSALES)})).sort_values("_o", kind="stable")


# ============================================================================
# Estilo comun
# ============================================================================
def _fill(c):
    return PatternFill("solid", fgColor=c)


def _font(color="FF1A1A2E", bold=False, size=9):
    return Font(name="Arial", color=color, bold=bold, size=size)


def _cel(ws, r, c, v=None, bg=None, fg="FF1A1A2E", bold=False, size=9, h=None, fmt=None, wrap=False):
    cell = ws.cell(r, c)
    cell.value = v
    if bg:
        cell.fill = _fill(bg)
    cell.font = _font(fg, bold, size)
    if h or wrap:
        cell.alignment = Alignment(horizontal=h, vertical="center", wrap_text=wrap)
    if fmt:
        cell.number_format = fmt
    return cell


def _hoja(wb, nombre, tab, titulo, subtitulo, ancho_banner, anchos):
    ws = wb.create_sheet(nombre)
    ws.sheet_properties.tabColor = tab
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    for col, w in anchos.items():
        ws.column_dimensions[col].width = w
    ultima = get_column_letter(ancho_banner)
    ws.merge_cells(f"A1:{ultima}1")
    _cel(ws, 1, 1, titulo, BANNER, "FFFFFFFF", True, 16, "left")
    for c in range(2, ancho_banner + 1):
        ws.cell(1, c).fill = _fill(BANNER)
    ws.merge_cells(f"A2:{ultima}2")
    _cel(ws, 2, 1, subtitulo, SUBT, SUBT_TXT, False, 9, "left")
    for c in range(2, ancho_banner + 1):
        ws.cell(2, c).fill = _fill(SUBT)
    ws.row_dimensions[1].height = 28
    return ws


def _encabezado(ws, r, c0, textos, bg):
    for i, t in enumerate(textos):
        b = bg[i] if isinstance(bg, list) else bg
        _cel(ws, r, c0 + i, t, b, "FFFFFFFF", True, 9, "center", wrap=True)
    ws.row_dimensions[r].height = 30


def _fila_total(ws, r, c0, c1, sumas, primera, ultima, fmts, etiqueta_col=None):
    for c in range(c0, c1 + 1):
        _cel(ws, r, c, None, TOTAL, "FFFFFFFF", True)
    _cel(ws, r, etiqueta_col or c0, "TOTAL", TOTAL, "FFFFFFFF", True)
    for c in sumas:
        L = get_column_letter(c)
        _cel(ws, r, c, f"=SUM({L}{primera}:{L}{ultima})", TOTAL, "FFFFFFFF", True, fmt=fmts.get(c))


def _detalle(ws, df, columnas, fmts, primera=5):
    """Tabla de detalle con fondo claro de la sucursal y la sucursal en negrita oscura."""
    r = primera
    for _, row in df.iterrows():
        oscuro, claro = COLOR_SUC.get(row["Suc"], ("FF5D6D7E", "FFF4F6F7"))
        for i, col in enumerate(columnas):
            v = row[col]
            if isinstance(v, float) and pd.isna(v):
                v = None
            _cel(ws, r, i + 1, v, claro, oscuro if i == 0 else "FF1A1A2E", i == 0, fmt=fmts.get(i + 1))
        r += 1
    return r - 1


def _barras(titulo, cats, series, colores, apilado, alto, ancho, etiquetas=False, sin_eje=False):
    ch = BarChart()
    ch.type = "bar"
    ch.title = titulo
    ch.style = 10
    ch.height, ch.width = alto, ancho
    if apilado:
        ch.grouping = "stacked"
        ch.overlap = 100
    for ref, color in zip(series, colores):
        ch.add_data(ref, titles_from_data=True)
        s = ch.series[-1]
        s.graphicalProperties.solidFill = color[2:]
        s.graphicalProperties.line.solidFill = color[2:]
    ch.set_categories(cats)
    ch.y_axis.majorGridlines = None
    ch.x_axis.scaling.orientation = "maxMin"      # el mayor arriba
    if etiquetas:
        ch.dataLabels = DataLabelList()
        ch.dataLabels.showVal = True
    if sin_eje:
        ch.y_axis.delete = True
    if not apilado:
        ch.legend = None
    return ch


# ============================================================================
# Productos
# ============================================================================
def _excluir_top(df):
    """Para los top 10: sin rubro DELIVERY, sin ENVIO..., sin cantidades <= 0."""
    return df[(df["Rubro"].str.upper() != "DELIVERY")
              & (~df["Producto"].str.upper().str.startswith("ENVIO"))
              & (df["Cantidad"] > 0)]


def hoja_productos(wb, d, fecha_txt):
    ws = _hoja(wb, "Productos", "FF784212", f"PRODUCTOS VENDIDOS  ·  {fecha_txt}",
               "Todas las unidades vendidas por sucursal, incluidas las que salieron dentro de promociones y "
               "sobreventas. Fuente: DF_DTW.", 13,
               {"A": 13, "B": 18, "C": 34, "D": 10, "E": 10, "F": 13, "G": 3, "H": 32,
                "I": 12, "J": 12, "K": 12, "L": 12, "M": 12, "N": 3})
    df = d["productos"].copy()
    df["Rubro"] = df["Rubro"].fillna("(sin rubro)")
    df = _orden_suc(df.sort_values(["Rubro", "Cantidad"], ascending=[True, False]))
    _encabezado(ws, 4, 1, ["Sucursal", "Rubro", "Producto", "Cantidad", "Kilos", "Venta"], BANNER)
    fmts = {4: F_CANT, 5: F_KG, 6: F_PESOS}
    ult = _detalle(ws, df, ["Suc", "Rubro", "Producto", "Cantidad", "Kilos", "Venta"], fmts)
    rt = ult + 1
    _fila_total(ws, rt, 1, 6, [4, 5, 6], 5, ult, fmts)
    ws.auto_filter.ref = f"A4:F{ult}"
    ws.freeze_panes = "A5"
    ws.conditional_formatting.add(f"D5:D{ult}", DataBarRule(start_type="min", end_type="max", color="C39B77"))
    rng = lambda L: f"${L}$5:${L}${ult}"  # rangos acotados al detalle

    # --- Top 10 de las 3 sucursales juntas (H4:M15)
    top = _excluir_top(df).groupby("Producto")["Cantidad"].sum().sort_values(ascending=False).head(10)
    ws.merge_cells("H4:M4")
    _cel(ws, 4, 8, "TOP 10 PRODUCTOS POR CANTIDAD  ·  LAS 3 SUCURSALES JUNTAS", BANNER, "FFFFFFFF", True, 9, "center")
    _encabezado(ws, 5, 8, ["Producto"] + SUCURSALES + ["Total", "% de las unidades"],
                [BANNER] + [COLOR_SUC[s][0] for s in SUCURSALES] + [TOTAL, TOTAL])
    for i, prod in enumerate(top.index):
        r = 6 + i
        bg = "FFFFFFFF" if i % 2 == 0 else "FFEAEDED"
        _cel(ws, r, 8, prod, bg)
        for j, s in enumerate(SUCURSALES):
            _cel(ws, r, 9 + j, f'=SUMIFS({rng("D")},{rng("A")},"{s}",{rng("C")},H{r})', bg, fmt=F_CANT)
        _cel(ws, r, 12, f"=SUM(I{r}:K{r})", bg, bold=True, fmt=F_CANT)
        _cel(ws, r, 13, f'=IFERROR(L{r}/SUM({rng("D")}),"")', bg, fmt=F_PCT)
    if len(top):
        u = 5 + len(top)
        ws.conditional_formatting.add(f"L6:L{u}", DataBarRule(start_type="min", end_type="max", color="5D6D7E"))
        cats = Reference(ws, min_col=8, min_row=6, max_row=u)
        series = [Reference(ws, min_col=9 + j, min_row=5, max_row=u) for j in range(3)]
        ch = _barras("Top 10 · aporte de cada sucursal", cats, series,
                     [COLOR_SUC[s][0] for s in SUCURSALES], True, 9, 17)
        ws.add_chart(ch, "O4")

    # --- Top 10 por sucursal: un bloque de 22 filas cada uno desde la fila 26
    for k, s in enumerate(SUCURSALES):
        r0 = 26 + 22 * k
        oscuro, claro = COLOR_SUC[s]
        t = (_excluir_top(df[df["Suc"] == s]).groupby("Producto")["Cantidad"].sum()
             .sort_values(ascending=False).head(10))
        ws.merge_cells(f"H{r0}:K{r0}")
        _cel(ws, r0, 8, f"TOP 10 PRODUCTOS POR CANTIDAD  ·  {s.upper()}", oscuro, "FFFFFFFF", True, 9, "center")
        _encabezado(ws, r0 + 1, 8, ["Producto", "Cantidad", "Venta", "% de las unidades"], oscuro)
        for i, prod in enumerate(t.index):
            r = r0 + 2 + i
            bg = "FFFFFFFF" if i % 2 == 0 else claro
            _cel(ws, r, 8, prod, bg)
            _cel(ws, r, 9, f'=SUMIFS({rng("D")},{rng("A")},"{s}",{rng("C")},H{r})', bg, fmt=F_CANT)
            _cel(ws, r, 10, f'=SUMIFS({rng("F")},{rng("A")},"{s}",{rng("C")},H{r})', bg, fmt=F_PESOS)
            _cel(ws, r, 11, f'=IFERROR(I{r}/SUMIF({rng("A")},"{s}",{rng("D")}),"")', bg, fmt=F_PCT)
        if len(t):
            u = r0 + 1 + len(t)
            ws.conditional_formatting.add(f"I{r0 + 2}:I{u}",
                                          DataBarRule(start_type="min", end_type="max", color=oscuro[2:]))
            ch = _barras(f"Top 10 · {s}", Reference(ws, min_col=8, min_row=r0 + 2, max_row=u),
                         [Reference(ws, min_col=9, min_row=r0 + 1, max_row=u)], [oscuro],
                         False, 9, 15, etiquetas=True, sin_eje=True)
            ws.add_chart(ch, f"O{r0}")
    return {"filas": len(df), "cantidad": float(df["Cantidad"].sum()), "kilos": float(df["Kilos"].sum()),
            "venta": float(df["Venta"].sum()), "top": top}


# ============================================================================
# Promociones
# ============================================================================
def hoja_promociones(wb, d, fecha_txt):
    ws = _hoja(wb, "Promociones", PROMO_COLOR, f"PROMOCIONES  ·  {fecha_txt}",
               "Qué productos salieron dentro de cada promoción y cuántas veces se usó cada una. Fuente: DF_DTW.",
               9 + len(SUCURSALES) + 1,
               {"A": 13, "B": 36, "C": 30, "D": 10, "E": 10, "F": 13, "G": 13, "H": 3, "I": 40,
                **{get_column_letter(10 + j): 12 for j in range(len(SUCURSALES))},
                get_column_letter(10 + len(SUCURSALES)): 11})
    df = _orden_suc(d["promo_det"].sort_values(["Promocion", "Cantidad"], ascending=[True, False]))
    _encabezado(ws, 4, 1, ["Sucursal", "Promoción", "Producto", "Cantidad", "Kilos", "Venta", "Descuento"],
                PROMO_COLOR)
    fmts = {4: F_CANT, 5: F_KG, 6: F_PESOS, 7: F_PESOS}
    if df.empty:
        _cel(ws, 5, 1, "Sin promociones en el día", bold=True)
        ult = 5
    else:
        ult = _detalle(ws, df, ["Suc", "Promocion", "Producto", "Cantidad", "Kilos", "Venta", "Descuento"], fmts)
        _fila_total(ws, ult + 1, 1, 7, [4, 5, 6, 7], 5, ult, fmts)
        ws.auto_filter.ref = f"A4:G{ult}"
    ws.freeze_panes = "A5"

    # --- Usos por promocion (pivot)
    usos = d["promo_usos"]
    piv = (usos.pivot_table(index="Promocion", columns="Suc", values="Usos", aggfunc="sum", fill_value=0)
           .reindex(columns=SUCURSALES, fill_value=0))
    piv["_t"] = piv.sum(axis=1)
    piv = piv[piv["_t"] > 0].sort_values("_t", ascending=False)
    cT = 10 + len(SUCURSALES)
    ws.merge_cells(start_row=4, start_column=9, end_row=4, end_column=cT)
    _cel(ws, 4, 9, "USOS POR PROMOCIÓN", PROMO_COLOR, "FFFFFFFF", True, 9, "center")
    _encabezado(ws, 5, 9, ["Promoción"] + SUCURSALES + ["Total"],
                [PROMO_COLOR] + [COLOR_SUC[s][0] for s in SUCURSALES] + [TOTAL])
    for i, (promo, row) in enumerate(piv.iterrows()):
        r = 6 + i
        bg = "FFEBF5FB" if i % 2 == 0 else "FFFFFFFF"
        _cel(ws, r, 9, promo, bg)
        for j, s in enumerate(SUCURSALES):
            _cel(ws, r, 10 + j, int(row[s]), bg, fmt=F_CANT)
        _cel(ws, r, cT, f"=SUM({get_column_letter(10)}{r}:{get_column_letter(cT - 1)}{r})", bg, bold=True,
             fmt=F_CANT)
    u = 5 + len(piv)
    rt = u + 1
    _fila_total(ws, rt, 9, cT, list(range(10, cT + 1)), 6, max(u, 6), {c: F_CANT for c in range(10, cT + 1)})
    _cel(ws, rt + 1, 9, "Usos = veces que se aplicó la promoción. Las promos de apps (PedidosYa / Apps) no "
                        "registran descuento en el POS.", fg="FF5D6D7E", size=8)
    if len(piv):
        ch = _barras("Usos por promoción", Reference(ws, min_col=9, min_row=6, max_row=u),
                     [Reference(ws, min_col=10 + j, min_row=5, max_row=u) for j in range(len(SUCURSALES))],
                     [COLOR_SUC[s][0] for s in SUCURSALES], True, max(7, 0.6 * len(piv) + 3), 20)
        ws.add_chart(ch, f"I{rt + 3}")
    return {"usos": {s: int(piv[s].sum()) for s in SUCURSALES}, "cantidad": float(df["Cantidad"].sum()),
            "venta": float(df["Venta"].sum()), "descuento": float(df["Descuento"].sum())}


# ============================================================================
# Sobreventas
# ============================================================================
def hoja_sobreventas(wb, d, fecha_txt, filas_suc):
    ws = _hoja(wb, "Sobreventas", "FF4A7C59", f"SOBREVENTAS  ·  {fecha_txt}",
               "Qué productos se llevaron como sobreventa y cómo convirtió cada sucursal. Ofrecidas y aceptadas "
               "salen de la hoja «Informe Diario».", 14,
               {"A": 13, "B": 24, "C": 30, "D": 10, "E": 10, "F": 13, "G": 13, "H": 3,
                "I": 30, "J": 11, "K": 11, "L": 11, "M": 12, "N": 10})
    df = _orden_suc(d["sv_det"].sort_values(["Sobreventa", "Cantidad"], ascending=[True, False]))
    _encabezado(ws, 4, 1, ["Sucursal", "Sobreventa", "Producto", "Cantidad", "Kilos", "Venta", "Descuento"],
                SV_COLOR)
    fmts = {4: F_CANT, 5: F_KG, 6: F_PESOS, 7: F_PESOS}
    if df.empty:
        _cel(ws, 5, 1, "Sin sobreventas aceptadas en el día", bold=True)
        ult = 5
    else:
        ult = _detalle(ws, df, ["Suc", "Sobreventa", "Producto", "Cantidad", "Kilos", "Venta", "Descuento"], fmts)
        _fila_total(ws, ult + 1, 1, 7, [4, 5, 6, 7], 5, ult, fmts)
        ws.auto_filter.ref = f"A4:G{ult}"
    ws.freeze_panes = "A5"

    # --- Ofrecidas / aceptadas, con formula al Informe Diario (columnas J y K)
    ws.merge_cells("I4:N4")
    _cel(ws, 4, 9, "OFRECIDAS Y ACEPTADAS  ·  según Informe Diario", SV_COLOR, "FFFFFFFF", True, 9, "center")
    _encabezado(ws, 5, 9, ["Sucursal", "Ofrecidas", "Aceptadas", "Rechazadas", "% aceptación", "Estándar"],
                SV_COLOR)
    r = 6
    for s in SUCURSALES:
        fr = filas_suc.get(s)
        if not fr:
            continue
        oscuro, claro = COLOR_SUC[s]
        _cel(ws, r, 9, f"='Informe Diario'!A{fr}", claro, oscuro, True)
        _cel(ws, r, 10, f"='Informe Diario'!J{fr}", claro, fmt=F_CANT)
        _cel(ws, r, 11, f"='Informe Diario'!K{fr}", claro, fmt=F_CANT)
        _cel(ws, r, 12, f"=J{r}-K{r}", claro, fmt=F_CANT)
        _cel(ws, r, 13, f'=IFERROR(K{r}/J{r},"")', claro, bold=True, h="center", fmt=F_PCT)
        _cel(ws, r, 14, UMBRAL_SV, claro, h="center", fmt=F_PCT)
        r += 1
    u = r - 1
    _fila_total(ws, r, 9, 14, [10, 11, 12], 6, u, {c: F_CANT for c in (10, 11, 12)})
    _cel(ws, r, 13, f'=IFERROR(K{r}/J{r},"")', TOTAL, "FFFFFFFF", True, h="center", fmt=F_PCT)
    _cel(ws, r, 14, UMBRAL_SV, TOTAL, "FFFFFFFF", True, h="center", fmt=F_PCT)
    ws.conditional_formatting.add(f"M6:M{r}", CellIsRule(
        operator="lessThan", formula=[str(UMBRAL_SV)], fill=_fill("FFC0392B"), font=_font("FFFFFFFF", True)))
    ws.conditional_formatting.add(f"M6:M{r}", CellIsRule(
        operator="greaterThanOrEqual", formula=[str(UMBRAL_SV)], fill=_fill("FF1E8449"),
        font=_font("FFFFFFFF", True)))
    r_tot_sv = r

    # --- Aceptadas por producto, 3 filas debajo del total
    r0 = r + 3
    ws.merge_cells(f"I{r0}:L{r0}")
    _cel(ws, r0, 9, "ACEPTADAS POR PRODUCTO", SV_COLOR, "FFFFFFFF", True, 9, "center")
    _encabezado(ws, r0 + 1, 9, ["Producto", "Cantidad", "Venta", "% de las unidades"], SV_COLOR)
    prods = df.groupby("Producto")["Cantidad"].sum().sort_values(ascending=False) if not df.empty else pd.Series(dtype=float)
    rd = lambda L: f"${L}$5:${L}${ult}"
    for i, prod in enumerate(prods.index):
        rr = r0 + 2 + i
        _cel(ws, rr, 9, prod, PALETA[i % len(PALETA)], "FFFFFFFF", True)
        _cel(ws, rr, 10, f'=SUMIF({rd("C")},I{rr},{rd("D")})', fmt=F_CANT)
        _cel(ws, rr, 11, f'=SUMIF({rd("C")},I{rr},{rd("F")})', fmt=F_PESOS)
        _cel(ws, rr, 12, f'=IFERROR(J{rr}/SUM({rd("D")}),"")', fmt=F_PCT)
    up = r0 + 1 + len(prods)
    _cel(ws, up + 1, 9, "Un ticket puede llevar más de una sobreventa: por eso la suma por producto puede superar "
                        "las aceptadas del Informe Diario.", fg="FF5D6D7E", size=8)

    # --- Graficos
    if u >= 6:
        ch = BarChart()
        ch.type = "col"
        ch.grouping = "stacked"
        ch.overlap = 100
        ch.title = "Aceptadas y rechazadas por sucursal"
        ch.height, ch.width = 8, 12
        for col, color in ((11, SV_COLOR), (12, "FFD5D8DC")):
            ch.add_data(Reference(ws, min_col=col, min_row=5, max_row=u), titles_from_data=True)
            ch.series[-1].graphicalProperties.solidFill = color[2:]
        ch.set_categories(Reference(ws, min_col=9, min_row=6, max_row=u))
        ch.y_axis.majorGridlines = None
        ws.add_chart(ch, f"I{up + 4}")
    if len(prods):
        dn = DoughnutChart()
        dn.title = "Sobreventas por producto"
        dn.height, dn.width = 8, 12
        dn.add_data(Reference(ws, min_col=10, min_row=r0 + 1, max_row=up), titles_from_data=True)
        dn.set_categories(Reference(ws, min_col=9, min_row=r0 + 2, max_row=up))
        for i in range(len(prods)):
            pt = DataPoint(idx=i)
            pt.graphicalProperties.solidFill = PALETA[i % len(PALETA)][2:]
            dn.series[0].dPt.append(pt)
        ws.add_chart(dn, f"L{up + 4}")
    return {"unidades": float(df["Cantidad"].sum()) if not df.empty else 0.0,
            "venta": float(df["Venta"].sum()) if not df.empty else 0.0,
            "fila_total_sv": r_tot_sv}


# ============================================================================
# Tickets
# ============================================================================
def hoja_tickets(wb, d, fecha):
    titulo = f"TICKETS DEL DÍA  ·  {DIAS[fecha.weekday()]} {fecha:%d/%m/%Y}"
    ws = _hoja(wb, "Tickets", "FF5D6D7E", titulo,
               "Un renglón por ticket: qué se llevó el cliente y qué pasó con la sobreventa. Usá los filtros "
               "(por ejemplo «Cajero») para ver a una persona.", 9,
               {"A": 13, "B": 13, "C": 8, "D": 11, "E": 12, "F": 18, "G": 70, "H": 13, "I": 24})
    df = d["tickets"].copy()
    # Cajero con el nombre como figura en el Informe Diario (alla va .title()).
    df["CajeroTxt"] = df["Cajero"].fillna("").str.strip().str.title()
    df["_o"] = df["Sucursal"].map({s: i for i, s in enumerate(SUCURSALES)})
    df = df.sort_values(["_o", "CajeroTxt", "FechaHora", "Ticket"])
    _encabezado(ws, 4, 1, ["Sucursal", "Cajero", "Hora", "Canal", "Importe", "Rubro principal",
                           "Producto llevado", "Sobreventa", "Producto llevado en SV"],
                [BANNER] * 7 + [SV_COLOR] * 2)
    estilo_sv = {"Aceptada": ("FF1E8449", "FFFFFFFF", True), "Rechazada": ("FFFADBD8", "FFC0392B", True),
                 "No ofrecida": ("FFEAEDED", "FF7F8C8D", False)}
    r = 5
    for _, t in df.iterrows():
        oscuro, claro = COLOR_SUC.get(t["Sucursal"], ("FF5D6D7E", "FFF4F6F7"))
        _cel(ws, r, 1, t["Sucursal"], claro, oscuro, True)
        _cel(ws, r, 2, t["CajeroTxt"], claro)
        _cel(ws, r, 3, f"{int(t['Hora'])} h" if pd.notna(t["Hora"]) else None, claro, h="center")
        _cel(ws, r, 4, t["Canal"], claro)
        _cel(ws, r, 5, float(t["Importe"]) if pd.notna(t["Importe"]) else None, claro, fmt=F_PESOS)
        _cel(ws, r, 6, t["RubroPrincipal"], claro)
        _cel(ws, r, 7, t["ProductoLlevado"] if pd.notna(t["ProductoLlevado"]) else None, claro)
        bg, fg, b = estilo_sv[t["Sobreventa"]]
        _cel(ws, r, 8, t["Sobreventa"], bg, fg, b, h="center")
        sv = t["ProductoLlevadoEnSV"] if pd.notna(t["ProductoLlevadoEnSV"]) else None
        _cel(ws, r, 9, sv, claro, bold=bool(sv))
        r += 1
    ult = r - 1
    for c in range(1, 10):
        _cel(ws, r, c, None, TOTAL, "FFFFFFFF", True)
    _cel(ws, r, 1, "TOTAL", TOTAL, "FFFFFFFF", True)
    # SUBTOTAL: al filtrar por un cajero, el total muestra solo lo de ese cajero.
    _cel(ws, r, 2, f"=SUBTOTAL(103,A5:A{ult})", TOTAL, "FFFFFFFF", True, fmt=F_CANT)
    _cel(ws, r, 5, f"=SUBTOTAL(109,E5:E{ult})", TOTAL, "FFFFFFFF", True, fmt=F_PESOS)
    _cel(ws, r + 2, 1, "Importe = total cobrado del ticket, ya neto de los descuentos de PedidosYa / Rappi: la "
                       "suma coincide con la venta del Informe Diario. En los pedidos de apps no se ofrece "
                       "sobreventa.", fg="FF5D6D7E", size=8)
    ws.auto_filter.ref = f"A4:I{ult}"
    ws.freeze_panes = "A5"
    return {"renglones": len(df), "importe": float(df["Importe"].sum()),
            "sv": dict(Counter(df["Sobreventa"]))}


# ============================================================================
# Efemerides
# ============================================================================
def _leer_csv(nombre):
    ruta = os.path.join(DIR_CALENDARIO, nombre)
    if not os.path.exists(ruta):
        return []
    with open(ruta, encoding="utf-8-sig", newline="") as f:
        filas = list(csv.DictReader(f, delimiter=";"))
    for x in filas:
        x["_fecha"] = datetime.strptime(x["Fecha"].strip(), "%Y-%m-%d").date()
    return filas


COLOR_EFEM = {"Escolar": "FFEBF5FB", "Familiar": "FFFEF9E7", "Ingresos": "FFEAFAF1", "Deportivo": "FFFDEDEC",
              "Electoral": "FFF4F6F7", "Comercial": "FFFEF5E7"}
TIPOS_FERIADO = ("Feriado", "Feriado provincial")


def hoja_efemerides(wb, fecha):
    MORADO = "FF7D3C98"
    ws = _hoja(wb, "Efemérides", MORADO, f"EFEMÉRIDES  ·  {DIAS[fecha.weekday()]} {fecha:%d/%m/%Y}",
               "Fechas especiales del día y del año: feriados, días no laborables, calendario escolar y otras "
               "fechas que pueden mover la venta. Se editan en calendario\\feriados.csv y efemerides.csv.", 7,
               {"A": 14, "B": 16, "C": 60, "D": 24, "E": 12, "F": 26, "G": 20})
    feriados, efem = _leer_csv("feriados.csv"), _leer_csv("efemerides.csv")

    ws.merge_cells("A4:E6")
    _cel(ws, 4, 1, "PRÓXIMAMENTE", "FFF4ECF7", MORADO, True, 28, "center")
    for rr in range(4, 7):
        for c in range(1, 6):
            ws.cell(rr, c).fill = _fill("FFF4ECF7")
    ws.merge_cells("A7:F7")
    _cel(ws, 7, 1, "En esta solapa se van a sumar las efemérides de cada día (fechas patrias, días especiales, "
                   "eventos de la zona). Las que ya están cargadas también las usa el modelo de venta esperada.",
         fg="FF5D6D7E", size=8)

    # --- EL DIA
    ws.merge_cells("A9:F9")
    _cel(ws, 9, 1, "EL DÍA", MORADO, "FFFFFFFF", True, 10)
    hoy_fer = [f for f in feriados if f["_fecha"] == fecha]
    hoy_ef = [e for e in efem if e["_fecha"] == fecha]
    prox = sorted((f for f in feriados if f["_fecha"] > fecha and f["Tipo"].strip() in TIPOS_FERIADO),
                  key=lambda f: f["_fecha"])
    prox = prox[0] if prox else None
    filas = [("Fecha", fecha, "dd/mm/yyyy", None),
             ("Día", DIAS_ES[fecha.weekday()], None, None)]
    for i, (k, v, fmt, _) in enumerate(filas):
        _cel(ws, 10 + i, 1, k, bold=True)
        _cel(ws, 10 + i, 2, v, fmt=fmt, h="left")
    _cel(ws, 12, 1, "¿Feriado?", bold=True)
    if hoy_fer:
        _cel(ws, 12, 2, "Sí", "FFC0392B", "FFFFFFFF", True, h="center")
        _cel(ws, 12, 3, " / ".join(f["Nombre"] for f in hoy_fer))
    else:
        _cel(ws, 12, 2, "No", "FF1E8449", "FFFFFFFF", True, h="center")
    _cel(ws, 13, 1, "Efeméride", bold=True)
    if hoy_ef:
        _cel(ws, 13, 2, "Sí", "FFD68910", "FFFFFFFF", True, h="center")
        _cel(ws, 13, 3, " / ".join(e["Nombre"] for e in hoy_ef))
    else:
        _cel(ws, 13, 2, "No", "FFEAEDED", "FF7F8C8D", h="center")
    _cel(ws, 14, 1, "Próximo feriado", bold=True)
    if prox:
        _cel(ws, 14, 2, prox["_fecha"], fmt="dd/mm/yyyy", h="left")
        _cel(ws, 14, 3, f"{DIAS_ES[prox['_fecha'].weekday()]}: {prox['Nombre']}")
    _cel(ws, 15, 1, "Faltan (días)", bold=True)
    _cel(ws, 15, 2, "=B14-B10" if prox else None, fmt="0", h="left")

    # --- CALENDARIO: desde el 1/1 del año del informe hasta el 1/5 del siguiente
    desde, hasta = date(fecha.year, 1, 1), date(fecha.year + 1, 5, 1)
    ws.merge_cells("A17:F17")
    _cel(ws, 17, 1, f"CALENDARIO {fecha.year}: FERIADOS, DÍAS NO LABORABLES Y FECHAS QUE PUEDEN AFECTAR LAS VENTAS",
         MORADO, "FFFFFFFF", True, 10)
    _encabezado(ws, 18, 1, ["Fecha", "Día", "Qué se celebra / qué pasa", "Tipo", "Faltan (días)", "Verificado"],
                MORADO)
    items = [(f["_fecha"], f["Nombre"], f["Tipo"].strip(), f.get("Verificado", ""), None) for f in feriados] + \
            [(e["_fecha"], e["Nombre"], f"Efeméride · {e['Tipo'].strip()}", e.get("Verificado", ""),
              e["Tipo"].strip()) for e in efem]
    items = sorted((x for x in items if desde <= x[0] <= hasta), key=lambda x: (x[0], x[2]))
    r = 19
    for fch, nombre, tipo, verif, tipo_ef in items:
        bg = COLOR_EFEM.get(tipo_ef) if tipo_ef else None
        if fch == fecha:
            bg = "FFF5B7B1"
        elif prox and fch == prox["_fecha"] and tipo in TIPOS_FERIADO:
            bg = "FFFDEBD0"
        fg = "FFA6ACAF" if fch < fecha else "FF1A1A2E"
        _cel(ws, r, 1, fch, bg, fg, fmt="dd/mm/yyyy", h="left")
        _cel(ws, r, 2, DIAS_ES[fch.weekday()], bg, fg)
        _cel(ws, r, 3, nombre, bg, fg)
        _cel(ws, r, 4, tipo, bg, fg)
        _cel(ws, r, 5, f"=A{r}-$B$10", bg, fg, fmt="0", h="center")
        verif = (verif or "").strip()
        if not verif or verif.lower() == "no":
            _cel(ws, r, 6, "A verificar", bg, "FFC0392B", True)
        else:
            _cel(ws, r, 6, verif, bg, fg)
        if prox and fch == prox["_fecha"] and tipo in TIPOS_FERIADO:
            _cel(ws, r, 7, "← próximo feriado", None, "FFD35400", True)
        r += 1
    if r > 19:
        ws.auto_filter.ref = f"A18:F{r - 1}"
    _cel(ws, r + 1, 1, "Fuente: calendario\\feriados.csv y efemerides.csv, que mantiene el equipo. Los feriados "
                       "trasladables van en el día en que se gozan. «A verificar» = cargado de memoria, confirmar "
                       "con el calendario oficial (y la DGCyE para las fechas escolares).", fg="FF5D6D7E", size=8)
    return {"feriado_hoy": bool(hoy_fer), "proximo": prox["_fecha"] if prox else None, "items": len(items)}


# ============================================================================
# Entrada
# ============================================================================
def filas_sucursal(ws_informe):
    """{sucursal: fila} de las filas de sucursal (B = "Todos") del Informe Diario."""
    out = {}
    for r in range(1, ws_informe.max_row + 1):
        if ws_informe.cell(r, 2).value == "Todos":
            out[str(ws_informe.cell(r, 1).value).strip()] = r
    return out


def agregar_hojas(wb, datos, fecha):
    """Agrega las 5 hojas al libro y devuelve los numeros para el log y la validacion."""
    if isinstance(fecha, datetime):
        fecha = fecha.date()
    fecha_txt = f"{fecha:%d/%m/%Y}"
    fs = filas_sucursal(wb["Informe Diario"])
    return {
        "productos": hoja_productos(wb, datos, fecha_txt),
        "promociones": hoja_promociones(wb, datos, fecha_txt),
        "sobreventas": hoja_sobreventas(wb, datos, fecha_txt, fs),
        "tickets": hoja_tickets(wb, datos, fecha),
        "efemerides": hoja_efemerides(wb, fecha),
    }
